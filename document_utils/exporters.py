import base64
import io
import os
import re
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF

from document_utils.formatting import sanitize_text


def _logo_bytes(
    logo_base64: str | None,
) -> bytes | None:

    if not logo_base64:
        return None

    try:
        if "," in logo_base64:
            logo_base64 = logo_base64.split(
                ",",
                1,
            )[1]

        return base64.b64decode(
            logo_base64,
        )

    except Exception:
        return None


def _set_docx_default_font(
    document: Document,
):
    style = document.styles["Normal"]

    style.font.name = "Times New Roman"
    style.font.size = Pt(11)


def _add_docx_footer(
    document: Document,
    footer_text: str,
):
    section = document.sections[0]

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = footer.add_run(
        footer_text,
    )

    run.font.name = "Times New Roman"
    run.font.size = Pt(8)


def format_docx(
    text: str,
    doc_type: str,
    company_name: str = "LegalEase",
    footer_text: str = "",
    logo_base64: str | None = None,
) -> bytes:

    document = Document()

    _set_docx_default_font(
        document,
    )

    section = document.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    logo = _logo_bytes(
        logo_base64,
    )

    if logo:
        try:
            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            paragraph.add_run().add_picture(
                io.BytesIO(logo),
                width=Inches(1.35),
            )

        except Exception:
            pass

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        doc_type.upper(),
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    brand = document.add_paragraph()

    brand.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    brand_run = brand.add_run(
        company_name,
    )

    brand_run.italic = True
    brand_run.font.size = Pt(9)

    for raw in sanitize_text(text).splitlines():

        line = raw.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(5)

        is_heading = (
            line.isupper()
            or bool(
                re.match(
                    r"^(SECTION|ARTICLE)\s+\d+",
                    line,
                    re.I,
                )
            )
        )

        if is_heading:

            run = paragraph.add_run(
                line,
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

        elif line.startswith(
            ("- ", "* "),
        ):

            paragraph.style = (
                document.styles["List Bullet"]
            )

            paragraph.add_run(
                line[2:].strip(),
            )

        else:

            paragraph.add_run(
                line,
            )

    if footer_text:
        _add_docx_footer(
            document,
            footer_text,
        )

    buffer = io.BytesIO()

    document.save(buffer)

    return buffer.getvalue()


class BrandedPDF(FPDF):

    def __init__(
        self,
        company_name: str,
        footer_text: str,
        logo_path: str | None = None,
    ):

        super().__init__()

        self.company_name = company_name
        self.footer_text = footer_text
        self.logo_path = logo_path

    def header(self):

        if (
            self.logo_path
            and os.path.exists(
                self.logo_path,
            )
        ):

            try:

                self.image(
                    self.logo_path,
                    x=95,
                    y=8,
                    w=20,
                )

                self.set_y(31)

            except Exception:

                self.set_y(10)

        else:

            self.set_y(10)

        self.set_x(self.l_margin)

        self.set_font(
            "Helvetica",
            "B",
            9,
        )

        self.cell(
            self.epw,
            5,
            self.company_name,
            align="C",
        )

        self.ln(7)

    def footer(self):

        self.set_y(-15)

        self.set_x(self.l_margin)

        self.set_font(
            "Helvetica",
            "",
            7,
        )

        self.multi_cell(
            self.epw,
            4,
            self.footer_text,
            align="C",
        )


def format_pdf(
    text: str,
    doc_type: str,
    company_name: str = "LegalEase",
    footer_text: str = "",
    logo_path: str | None = None,
) -> bytes:

    pdf = BrandedPDF(
        company_name,
        footer_text,
        logo_path,
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    pdf.set_x(pdf.l_margin)

    pdf.set_font(
        "Helvetica",
        "B",
        15,
    )

    pdf.cell(
        pdf.epw,
        10,
        doc_type.upper(),
        align="C",
    )

    pdf.ln(12)

    for raw in sanitize_text(text).splitlines():

        line = raw.strip()

        if not line:
            pdf.ln(3)
            continue

        is_heading = (
            line.isupper()
            or bool(
                re.match(
                    r"^(SECTION|ARTICLE)\s+\d+",
                    line,
                    re.I,
                )
            )
        )

        if is_heading:

            pdf.set_font(
                "Helvetica",
                "B",
                11,
            )

            pdf.set_x(pdf.l_margin)

            pdf.multi_cell(
                pdf.epw,
                6,
                line,
            )

        elif line.startswith(
            ("- ", "* "),
        ):

            pdf.set_font(
                "Helvetica",
                "",
                10,
            )

            pdf.set_x(pdf.l_margin)

            pdf.multi_cell(
                pdf.epw,
                5,
                "- " + line[2:].strip(),
            )

        else:

            pdf.set_font(
                "Helvetica",
                "",
                10,
            )

            pdf.set_x(pdf.l_margin)

            pdf.multi_cell(
                pdf.epw,
                5,
                line,
            )

    return bytes(
        pdf.output()
    )


def format_txt(
    text: str,
) -> bytes:

    return sanitize_text(
        text,
    ).encode("utf-8")